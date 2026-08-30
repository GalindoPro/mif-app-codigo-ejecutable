#!/usr/bin/env python3
import os
import sys
import json
import openpyxl

def clean_float(val, fallback=0.0):
    if val is None:
        return fallback
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("Q", "").replace(",", "").strip()
    while ".." in s:
        s = s.replace("..", ".")
    try:
        return float(s)
    except:
        return fallback

def extract():
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))

    # 1. Socios y Aportaciones
    path_apor = os.path.join(base_dir, "caja/APORTACIONES 31-08-26.xlsx")
    socios = []
    if os.path.exists(path_apor):
        wb_apor = openpyxl.load_workbook(path_apor, data_only=True)
        ws_apor = wb_apor['MES DE ENERO'] if 'MES DE ENERO' in wb_apor.sheetnames else wb_apor.active
        for r in range(9, ws_apor.max_row + 1):
            nombre = ws_apor.cell(r, 5).value
            if not nombre or not str(nombre).strip() or str(nombre).strip().startswith('='):
                continue
            asoc_no = ws_apor.cell(r, 1).value
            dpi = ws_apor.cell(r, 6).value
            edad = ws_apor.cell(r, 7).value
            genero = ws_apor.cell(r, 8).value
            aportacion = ws_apor.cell(r, 9).value
            cuenta = ws_apor.cell(r, 13).value
            fecha = ws_apor.cell(r, 2).value
            recibo = ws_apor.cell(r, 3).value
            
            clean_dpi = str(dpi).strip().replace(" ", "").replace("-", "") if dpi else None
            clean_genero = "M"
            if genero and str(genero).strip().upper() in ["F", "FEMENINO", "MUJER"]:
                clean_genero = "F"
            elif genero and str(genero).strip().upper() in ["M", "MASCULINO", "HOMBRE"]:
                clean_genero = "M"

            socios.append({
                "numero_asociado": f"CHAJ-{str(asoc_no).strip().zfill(4)}" if asoc_no else f"CHAJ-{str(len(socios)+1).zfill(4)}",
                "nombre": str(nombre).strip().upper(),
                "dpi": clean_dpi,
                "edad": int(edad) if edad and str(edad).isdigit() else None,
                "genero": clean_genero,
                "monto_aportacion": clean_float(aportacion, 100.0),
                "numero_cuenta": str(cuenta).strip() if cuenta else f"APOR-{len(socios)+1}",
                "fecha_ingreso": str(fecha).split(" ")[0] if fecha else "2026-01-02",
                "numero_recibo": str(recibo).strip() if recibo else None,
            })

    # 2. Préstamos
    path_kardex = os.path.join(base_dir, "promotor/KARDEX PRESTAMOS 01-07-26 AL 31-07-26 promotor 2.xlsx")
    prestamos = []
    if os.path.exists(path_kardex):
        wb_kardex = openpyxl.load_workbook(path_kardex, data_only=True)
        for sheet_name, tipo in [("HIPOTECARIO", "HIPOTECARIO"), ("FIDUCIARIO", "FIDUCIARIO")]:
            if sheet_name not in wb_kardex.sheetnames:
                continue
            ws = wb_kardex[sheet_name]
            for r in range(7, ws.max_row + 1):
                nombre = ws.cell(r, 2).value
                if not nombre or not str(nombre).strip() or "TOTAL" in str(nombre).upper():
                    continue
                asoc_no = ws.cell(r, 1).value
                doc = ws.cell(r, 4).value

                if tipo == "HIPOTECARIO":
                    ubi = ws.cell(r, 5).value or ws.cell(r, 6).value
                    fiador = ws.cell(r, 7).value
                    plazo_raw = str(ws.cell(r, 9).value or "")
                    f_pres = ws.cell(r, 10).value
                    f_venc = ws.cell(r, 11).value
                    monto_aprobado = ws.cell(r, 12).value
                    saldo_capital = ws.cell(r, 13).value
                else: # FIDUCIARIO
                    ubi = ws.cell(r, 5).value
                    fiador = ws.cell(r, 6).value
                    plazo_raw = str(ws.cell(r, 8).value or "")
                    f_pres = ws.cell(r, 9).value
                    f_venc = ws.cell(r, 10).value
                    monto_aprobado = ws.cell(r, 11).value
                    saldo_capital = ws.cell(r, 12).value

                plazo_meses = 120 if "10" in plazo_raw else (60 if "5" in plazo_raw else (36 if "3" in plazo_raw else (24 if "2" in plazo_raw else 12)))
                
                monto_num = clean_float(monto_aprobado, 10000.0)
                saldo_num = clean_float(saldo_capital, monto_num) if saldo_capital is not None else monto_num
                
                f_des_str = str(f_pres).split(" ")[0] if f_pres else "2026-01-01"
                f_venc_str = str(f_venc).split(" ")[0] if f_venc else "2027-01-01"

                cod = str(asoc_no).strip() if asoc_no else f"PREST-{tipo[:3]}-{len(prestamos)+1}"

                prestamos.append({
                    "codigo": cod,
                    "socio_nombre": str(nombre).strip().upper(),
                    "tipo": tipo,
                    "doc_desembolso": str(doc).strip() if doc else None,
                    "ubicacion_garantia": str(ubi).strip().upper() if ubi else None,
                    "fiador": str(fiador).strip().upper() if fiador else None,
                    "plazo_meses": plazo_meses,
                    "fecha_desembolso": f_des_str,
                    "fecha_vencimiento": f_venc_str,
                    "monto_aprobado": monto_num,
                    "saldo_capital": saldo_num,
                })

    # 3. Plazo Fijo
    path_pf = os.path.join(base_dir, "caja/KARDEX AHORRO PF 2026-08.xlsx")
    plazos_fijos = []
    if os.path.exists(path_pf):
        wb_pf = openpyxl.load_workbook(path_pf, data_only=True)
        ws_pf = wb_pf["Hoja1"] if "Hoja1" in wb_pf.sheetnames else wb_pf.active
        for r in range(10, ws_pf.max_row + 1):
            nombre = ws_pf.cell(r, 2).value
            cert = ws_pf.cell(r, 7).value
            deposito = ws_pf.cell(r, 10).value
            if not nombre or not cert or not deposito:
                continue
            if "TOTAL" in str(nombre).upper() or str(nombre).strip().startswith("="):
                continue
            cuenta = ws_pf.cell(r, 1).value
            plazo = ws_pf.cell(r, 8).value
            f_in = ws_pf.cell(r, 4).value
            f_ret = ws_pf.cell(r, 17).value
            rec_ret = ws_pf.cell(r, 18).value
            egreso = ws_pf.cell(r, 20).value
            
            f_in_str = str(f_in).split(" ")[0] if f_in else "2025-01-01"
            f_ret_str = str(f_ret).split(" ")[0] if f_ret else None
            monto_dep_num = clean_float(deposito, 1000.0)
            monto_liq_num = clean_float(egreso, monto_dep_num if f_ret_str else None)
            
            plazos_fijos.append({
                "socio_nombre": str(nombre).strip().upper(),
                "numero_cuenta": str(cuenta).strip() if cuenta else f"PF-{str(cert).strip()}",
                "numero_certificacion": str(cert).strip(),
                "plazo_meses": int(plazo) if plazo and str(plazo).isdigit() else 12,
                "monto_deposito": monto_dep_num,
                "fecha_inicio": f_in_str,
                "fecha_retiro": f_ret_str,
                "recibo_retiro": str(rec_ret).strip() if rec_ret else None,
                "monto_liquidado": monto_liq_num,
                "estado": "LIQUIDADO" if (f_ret_str or rec_ret) else "ACTIVO"
            })

    output = {
        "socios": socios,
        "prestamos": prestamos,
        "plazos_fijos": plazos_fijos
    }
    print(json.dumps(output))

if __name__ == "__main__":
    extract()
